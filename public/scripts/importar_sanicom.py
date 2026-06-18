#!/usr/bin/env python3
"""
Importador de Planilla Sanicom (Carlos)
========================================
Lee el archivo Excel con openpyxl para conservar los colores de celda.

Estructura exacta de la planilla:
  Fila 5  → encabezados
  Fila 9+ → datos (filas 6-8 son cabeceras visuales / espaciado)

  Col  1  Diatermia       ┐
  Col  2  O.Choque        │
  Col  3  ECO             │ Columnas de equipos
  Col  4  Magneto         │ Color de la celda indica:
  Col  5  Láser           │   FF00FF00 (verde)  → equipo que TIENE
  Col  6  Radio           │   FF00FFFF (celeste) → equipo con INTERÉS
  Col  7  Otros           │   FFFFFFFF (blanco)  → ignorar
  Col  8  Otros 2         │
  Col  9  Otros 3         ┘
  Col 12  NOMBRE
  Col 13  DIRECCION
  Col 14  POB.
  Col 15  PROV.
  Col 16  C.P
  Col 17  TELEFONO
  Col 18  E-MAIL
  Col 19  WEB
  Col 20  OBSERVACIONES

Uso:
  python importar_sanicom.py planilla.xlsx
  python importar_sanicom.py planilla.xlsx --verbose
  python importar_sanicom.py planilla.xlsx -o salida.json

Requiere: pip install openpyxl
"""

import json
import sys
import argparse
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("\nERROR: Necesitas instalar openpyxl:")
    print("  pip install openpyxl\n")
    sys.exit(1)


# ═══════════════════════════════════════════════════════════════════════════════
# ESTRUCTURA EXACTA DE LA PLANILLA DE CARLOS
# ═══════════════════════════════════════════════════════════════════════════════

FILA_CABECERA  = 5   # Fila con los encabezados
FILA_DATOS     = 9   # Primera fila con datos de clientes

# Columnas de equipos: índice (1-based) → nombre del equipo
COLS_EQUIPOS = {
    1:  'Diatermia',
    2:  'O.Choque',
    3:  'ECO',
    4:  'Magneto',
    5:  'Láser',
    6:  'Radio',
    7:  'Otros',
    8:  'Otros 2',
    9:  'Otros 3',
}

# Columnas de datos del cliente: índice (1-based)
COL_NOMBRE        = 12
COL_DIRECCION     = 13
COL_CIUDAD        = 14   # POB.
COL_PROVINCIA     = 15   # PROV.
COL_CP            = 16   # C.P
COL_TELEFONO      = 17
COL_EMAIL         = 18
COL_WEB           = 19
COL_OBSERVACIONES = 20

# Valores por defecto para todos los clientes importados
ESPECIALIDAD_DEFECTO = 'Fisioterapia'
TIPO_DEFECTO         = 'Clínica'
ESTADO_DEFECTO       = 'Activo'


# ═══════════════════════════════════════════════════════════════════════════════
# COLORES EXACTOS (ARGB en formato openpyxl: 8 caracteres hex AARRGGBB)
# ═══════════════════════════════════════════════════════════════════════════════

COLOR_VERDE   = 'FF00FF00'   # Verde puro  → equipo que TIENE
COLOR_CELESTE = 'FF00FFFF'   # Cyan puro   → equipo con INTERÉS
COLOR_BLANCO  = 'FFFFFFFF'   # Blanco      → ignorar
COLOR_TRANS   = '00000000'   # Transparente → ignorar


def get_cell_color_type(cell):
    """
    Retorna 'verde', 'celeste' o None según el color ARGB exacto del relleno.
    Comprueba fgColor (patrón sólido en Excel = fgColor tiene el color real).
    """
    try:
        fill = cell.fill
        if fill is None or fill.fill_type in (None, 'none'):
            return None

        argb = None

        # En Excel, los rellenos sólidos guardan el color en fgColor
        fgc = fill.fgColor
        if fgc and fgc.type == 'rgb':
            argb = fgc.rgb.upper()

        # Si fgColor es transparente, probar bgColor
        if not argb or argb in (COLOR_BLANCO, COLOR_TRANS):
            bgc = fill.bgColor
            if bgc and bgc.type == 'rgb':
                argb = bgc.rgb.upper()

        if not argb or argb in (COLOR_BLANCO, COLOR_TRANS):
            return None

        if argb == COLOR_VERDE:
            return 'verde'
        if argb == COLOR_CELESTE:
            return 'celeste'

    except Exception:
        pass

    return None


# ═══════════════════════════════════════════════════════════════════════════════
# PROCESADO
# ═══════════════════════════════════════════════════════════════════════════════

def normalizar(val):
    """Devuelve el valor como string limpio, o '' si es None/vacío."""
    if val is None:
        return ''
    s = str(val).strip()
    # Eliminar ".0" que Excel añade a números (ej. CP "41001.0" → "41001")
    if s.endswith('.0') and s[:-2].isdigit():
        s = s[:-2]
    return s


def cell_val(ws, row, col):
    return normalizar(ws.cell(row=row, column=col).value)


def procesar_hoja(ws, nombre_hoja, verbose=False):
    """Procesa una hoja y devuelve lista de clientes."""
    clientes = []

    for row_idx in range(FILA_DATOS, ws.max_row + 1):
        nombre = cell_val(ws, row_idx, COL_NOMBRE)
        if not nombre:
            continue  # Fila sin nombre → omitir

        provincia_celda = cell_val(ws, row_idx, COL_PROVINCIA)
        cliente = {
            'nombre':       nombre,
            'direccion':    cell_val(ws, row_idx, COL_DIRECCION),
            'ciudad':       cell_val(ws, row_idx, COL_CIUDAD) or nombre_hoja,
            'provincia':    provincia_celda or nombre_hoja,
            'cp':           cell_val(ws, row_idx, COL_CP),
            'telefono':     cell_val(ws, row_idx, COL_TELEFONO),
            'email':        cell_val(ws, row_idx, COL_EMAIL),
            'web':          cell_val(ws, row_idx, COL_WEB),
            'especialidad': ESPECIALIDAD_DEFECTO,
            'tipo':         TIPO_DEFECTO,
            'estado':       ESTADO_DEFECTO,
            'equiposInstalados': [],
            'equiposInteres':    [],
        }

        obs = cell_val(ws, row_idx, COL_OBSERVACIONES)
        if obs:
            cliente['notas'] = obs

        # Leer columnas de equipos por color
        tiene = []
        interes = []
        for col_idx, eq_nombre in COLS_EQUIPOS.items():
            cell = ws.cell(row=row_idx, column=col_idx)
            color = get_cell_color_type(cell)

            if color not in ('verde', 'celeste'):
                continue

            val_celda = normalizar(cell.value)
            # "ECO: V. Esaote mylab X6"  o solo "ECO" si la celda está vacía
            nombre_equipo = f"{eq_nombre}: {val_celda}" if val_celda else eq_nombre

            if color == 'verde':
                cliente['equiposInstalados'].append({
                    'id':              f"imp_{row_idx}_{col_idx}",
                    'nombre':          nombre_equipo,
                    'marca':           '',
                    'modelo':          '',
                    'nSerie':          '',
                    'anioInstalacion': '',
                    'distribuidor':    '',
                    'estado':          'Operativo',
                })
                tiene.append(nombre_equipo)

            elif color == 'celeste':
                cliente['equiposInteres'].append({
                    'id':     f"imp_int_{row_idx}_{col_idx}",
                    'nombre': nombre_equipo,
                })
                interes.append(nombre_equipo)

        clientes.append(cliente)

        if verbose and (tiene or interes):
            print(f"    [{nombre[:35]:<35}]  "
                  f"Tiene: {tiene or '-'}  |  Interés: {interes or '-'}")

    return clientes


def main():
    parser = argparse.ArgumentParser(description='Importador de planilla Sanicom')
    parser.add_argument('archivo', help='Ruta al archivo Excel (.xlsx)')
    parser.add_argument('-o', '--output', default=None,
                        help='Archivo JSON de salida (default: clientes_sanicom.json junto al Excel)')
    parser.add_argument('--verbose', '-v', action='store_true',
                        help='Muestra equipos detectados por fila')
    args = parser.parse_args()

    archivo = Path(args.archivo)
    if not archivo.exists():
        print(f"\nERROR: No se encuentra el archivo: {archivo}\n")
        sys.exit(1)

    print(f"\n📂 Leyendo: {archivo.name}")
    wb = openpyxl.load_workbook(archivo, data_only=True)
    print(f"   Hojas encontradas ({len(wb.sheetnames)}): {', '.join(wb.sheetnames)}")
    print(f"   Encabezados en fila {FILA_CABECERA}, datos desde fila {FILA_DATOS}\n")

    todos_clientes = []

    print(f"⚙️  Procesando...\n")
    for nombre_hoja in wb.sheetnames:
        ws = wb[nombre_hoja]
        clientes = procesar_hoja(ws, nombre_hoja, verbose=args.verbose)
        todos_clientes.extend(clientes)

        n_verde   = sum(1 for c in clientes if c['equiposInstalados'])
        n_celeste = sum(1 for c in clientes if c['equiposInteres'])
        print(f"  ✔ {nombre_hoja:<22} {len(clientes):3d} clientes  "
              f"(verde: {n_verde:3d}, celeste: {n_celeste:3d})")

    salida = Path(args.output) if args.output else archivo.parent / 'clientes_sanicom.json'
    with open(salida, 'w', encoding='utf-8') as f:
        json.dump(todos_clientes, f, ensure_ascii=False, indent=2)

    total          = len(todos_clientes)
    total_equipos  = sum(len(c['equiposInstalados']) for c in todos_clientes)
    total_interes  = sum(len(c['equiposInteres'])    for c in todos_clientes)

    print(f"\n{'─'*50}")
    print(f"✅ Total clientes:         {total}")
    print(f"   Equipos que tienen:    {total_equipos}  (celdas verdes  FF00FF00)")
    print(f"   Equipos con interés:   {total_interes}  (celdas celeste FF00FFFF)")
    print(f"\n📄 JSON guardado en: {salida}")
    print(f"\n👉 Sube ese JSON desde el CRM:")
    print(f"   Módulo Clientes → 'Importar planilla Sanicom' → selecciona {salida.name}\n")


if __name__ == '__main__':
    main()
